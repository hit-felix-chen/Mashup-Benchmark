#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import math
import re
import statistics
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_RESULTS_DIR = ROOT / "eval_results"

SPECIFIED_METRICS = [
    ("event", "事件型", "EC", "Event Coverage", "事件覆盖率"),
    ("event", "事件型", "KMS", "Key Moment Salience", "片段显著性"),
    ("character", "人物型", "PPR", "Protagonist Presence Rate", "主体出镜率"),
    ("character", "人物型", "PPM", "Protagonist Prominence", "主体显著性"),
    ("emotion", "情绪型", "ES", "Emotional Specificity", "情绪特异性"),
    ("emotion", "情绪型", "FBAE", "Facial / Body Affect Evidence", "面部/肢体情绪证据"),
    ("narrative", "叙事型", "NSC", "Narrative Structure Completeness", "叙事完整性"),
    ("narrative", "叙事型", "TLO", "Temporal / Logical Order", "时间合理性"),
]


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    rows = []
    with path.open("r", encoding="utf-8") as f:
        for line in f:
            if line.strip():
                rows.append(json.loads(line))
    return rows


def load_summary(result_dir: Path) -> dict[str, Any]:
    path = result_dir / "specified_metric_summary.json"
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def mean_or_none(values: list[float]) -> float | None:
    if not values:
        return None
    return float(statistics.fmean(values))


def normalized_score_key(value: Any) -> str:
    return str(value).strip().strip("'\"")


def to_float_score(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        score = float(value)
    except (TypeError, ValueError):
        return None
    return score if math.isfinite(score) else None


def score_value(row: dict[str, Any], metric: str) -> float | None:
    scores = row.get("scores") or {}
    normalized_scores = {normalized_score_key(key): value for key, value in scores.items()}
    return to_float_score(normalized_scores.get(metric))


def build_table(records: list[dict[str, Any]]) -> list[list[str | float | None]]:
    success_rows = [row for row in records if row.get("status") == "success"]
    table: list[list[str | float | None]] = []
    for task_type, type_zh, metric, _name_en, _name_zh in SPECIFIED_METRICS:
        values = [
            value
            for row in success_rows
            if row.get("task_type") == task_type
            if (value := score_value(row, metric)) is not None
        ]
        table.append([type_zh, metric, mean_or_none(values)])
    return table


def slugify(value: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9._-]+", "_", value.strip()).strip("_")
    return slug or "model"


def resolve_result_dir(results_dir: Path, specified_metrics_id: str) -> Path:
    result_dir = results_dir / specified_metrics_id
    if not result_dir.exists():
        raise FileNotFoundError(f"Specified metrics result folder not found: {result_dir}")
    return result_dir


def style_workbook(ws, *, num_rows: int) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    type_fill = PatternFill("solid", fgColor="EEF5FF")
    thin = Side(style="thin", color="B7C9D8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in ws.iter_rows(min_row=1, max_row=num_rows, min_col=1, max_col=3):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for row_idx in range(2, num_rows + 1):
        ws.cell(row=row_idx, column=1).fill = type_fill
        ws.cell(row=row_idx, column=1).font = Font(bold=True)
        ws.cell(row=row_idx, column=3).number_format = "0.00"

    widths = [14, 48, 16]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 22
    for row_idx in range(2, num_rows + 1):
        ws.row_dimensions[row_idx].height = 36


def write_xlsx(path: Path, model_name: str, table: list[list[str | float | None]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Specified Metrics"
    ws.append(["类型", "指标", model_name])
    for row in table:
        ws.append(row)
    style_workbook(ws, num_rows=len(table) + 1)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Export specified metrics result summary to Excel.")
    parser.add_argument("--model-name", required=True, help="Model or method name used as the third column header.")
    parser.add_argument("--specified-metrics-id", required=True, help="Specified metrics result folder name under eval_results/.")
    parser.add_argument("--results-dir", default=str(DEFAULT_RESULTS_DIR), help="Directory containing result folders.")
    parser.add_argument(
        "--output",
        default=None,
        help="Output .xlsx path. Defaults to eval_results/<specified_metrics_id>/<specified_metrics_id>_<model>_specified_metrics.xlsx.",
    )
    args = parser.parse_args()

    results_dir = Path(args.results_dir)
    if not results_dir.is_absolute():
        results_dir = ROOT / results_dir
    result_dir = resolve_result_dir(results_dir, args.specified_metrics_id)
    scores_path = result_dir / "specified_metric_scores.jsonl"
    if not scores_path.exists():
        raise FileNotFoundError(f"Specified metric scores file not found: {scores_path}")

    summary = load_summary(result_dir)
    evaluated_at = str(summary.get("created_at") or datetime.fromtimestamp(scores_path.stat().st_mtime, UTC).isoformat())
    table = build_table(load_jsonl(scores_path))

    output_path = (
        Path(args.output)
        if args.output
        else result_dir / f"{args.specified_metrics_id}_{slugify(args.model_name)}_specified_metrics.xlsx"
    )
    if not output_path.is_absolute():
        output_path = ROOT / output_path
    write_xlsx(output_path, args.model_name, table)
    print(f"evaluated_at={evaluated_at}")
    print(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
