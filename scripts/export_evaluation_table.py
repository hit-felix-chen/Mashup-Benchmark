#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import math
import re
import statistics
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
TASK_FILE = ROOT / "data" / "tasks" / "mashup_benchmark.jsonl"
DEFAULT_RESULTS_DIR = ROOT / "eval_results"

TASK_TYPES = [
    ("event", "事件型"),
    ("character", "人物型"),
    ("emotion", "情绪型"),
    ("narrative", "叙事型"),
]
METRICS = ["IF", "BCS", "AEC", "VQ", "TC", "NC", "Quality"]
HEADERS = ["评测指标", "事件型", "人物型", "情绪型", "叙事型", "总平均分"]


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    rows = []
    with path.open("r", encoding="utf-8") as f:
        for line in f:
            if line.strip():
                rows.append(json.loads(line))
    return rows


def load_tasks() -> dict[str, dict[str, Any]]:
    return {row["id"]: row for row in load_jsonl(TASK_FILE)}


def load_summary(eval_dir: Path) -> dict[str, Any]:
    path = eval_dir / "summary.json"
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def normalized_score_key(value: Any) -> str:
    return str(value).strip().strip("'\"")


def to_float_score(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        score = float(value)
    except (TypeError, ValueError):
        return None
    return score if math.isfinite(score) else None


def score_value(row: dict[str, Any], metric: str) -> float | None:
    scores = row.get("scores") or {}
    normalized_scores = {normalized_score_key(key): value for key, value in scores.items()}
    return to_float_score(normalized_scores.get(metric))


def mean_or_none(values: list[float]) -> float | None:
    if not values:
        return None
    return float(statistics.fmean(values))


def build_table(records: list[dict[str, Any]], tasks: dict[str, dict[str, Any]]) -> list[list[str | float | None]]:
    rows_by_type: dict[str, list[dict[str, Any]]] = {task_type: [] for task_type, _ in TASK_TYPES}
    success_rows = []
    for row in records:
        if row.get("status") != "success":
            continue
        task = tasks.get(row.get("task_id"))
        if not task:
            continue
        task_type = task.get("task", {}).get("type")
        if task_type not in rows_by_type:
            continue
        rows_by_type[task_type].append(row)
        success_rows.append(row)

    table = []
    for metric in METRICS:
        metric_row: list[str | float | None] = [metric]
        for task_type, _ in TASK_TYPES:
            values = [value for row in rows_by_type[task_type] if (value := score_value(row, metric)) is not None]
            metric_row.append(mean_or_none(values))
        all_values = [value for row in success_rows if (value := score_value(row, metric)) is not None]
        metric_row.append(mean_or_none(all_values))
        table.append(metric_row)
    return table


def slugify(value: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9._-]+", "_", value.strip()).strip("_")
    return slug or "model"


def resolve_eval_dir(results_dir: Path, eval_id: str) -> Path:
    eval_dir = results_dir / eval_id
    if not eval_dir.exists():
        raise FileNotFoundError(f"Evaluation folder not found: {eval_dir}")
    return eval_dir


def style_workbook(ws) -> None:
    title_fill = PatternFill("solid", fgColor="E7EEF8")
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="B7C9D8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("A1:F1")
    ws.merge_cells("A2:F2")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A2"].font = Font(italic=True)
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=1, max_row=10, min_col=1, max_col=6):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")

    for cell in ws[3]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for row_idx in range(4, 11):
        ws.cell(row=row_idx, column=1).font = Font(bold=True)
        for col_idx in range(2, 7):
            ws.cell(row=row_idx, column=col_idx).number_format = "0.00"

    widths = [16, 14, 14, 14, 14, 14]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 20


def write_xlsx(path: Path, model_name: str, evaluated_at: str, table: list[list[str | float | None]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    ws["A1"] = model_name
    ws["A2"] = f"评测时间：{evaluated_at}"
    ws.append(HEADERS)
    for row in table:
        ws.append(row)

    style_workbook(ws)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Export an evaluation result summary table to Excel.")
    parser.add_argument("--model-name", required=True, help="Model or method name shown in the first row.")
    parser.add_argument("--eval-id", required=True, help="Evaluation result folder name under eval_results/.")
    parser.add_argument("--results-dir", default=str(DEFAULT_RESULTS_DIR), help="Directory containing evaluation result folders.")
    parser.add_argument("--output", default=None, help="Output .xlsx path. Defaults to eval_results/<eval_id>/<eval_id>_<model>_summary.xlsx.")
    args = parser.parse_args()

    results_dir = Path(args.results_dir)
    if not results_dir.is_absolute():
        results_dir = ROOT / results_dir
    eval_dir = resolve_eval_dir(results_dir, args.eval_id)
    scores_path = eval_dir / "evaluation_scores.jsonl"
    if not scores_path.exists():
        raise FileNotFoundError(f"Evaluation scores file not found: {scores_path}")

    summary = load_summary(eval_dir)
    evaluated_at = str(summary.get("created_at") or datetime.fromtimestamp(scores_path.stat().st_mtime, UTC).isoformat())
    table = build_table(load_jsonl(scores_path), load_tasks())

    output_path = Path(args.output) if args.output else eval_dir / f"{args.eval_id}_{slugify(args.model_name)}_summary.xlsx"
    if not output_path.is_absolute():
        output_path = ROOT / output_path
    write_xlsx(output_path, args.model_name, evaluated_at, table)
    print(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
