#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import math
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
TASK_FILE = ROOT / "data" / "tasks" / "mashup_benchmark.jsonl"
DEFAULT_RESULTS_DIR = ROOT / "eval_results"

GENERAL_METRICS = ["IF", "BCS", "AEC", "VQ", "TC", "NC", "Quality"]
SPECIFIED_METRICS_BY_TYPE = {
    "event": ["EC", "KMS"],
    "character": ["PPR", "PPM"],
    "emotion": ["ES", "FBAE"],
    "narrative": ["NSC", "TLO"],
}
SPECIFIED_METRICS = ["EC", "KMS", "PPR", "PPM", "ES", "FBAE", "NSC", "TLO"]
TASK_TYPE_ZH = {
    "event": "事件型",
    "character": "人物型",
    "emotion": "情绪型",
    "narrative": "叙事型",
}


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    rows = []
    with path.open("r", encoding="utf-8") as f:
        for line in f:
            if line.strip():
                rows.append(json.loads(line))
    return rows


def load_tasks() -> dict[str, dict[str, Any]]:
    return {row["id"]: row for row in load_jsonl(TASK_FILE)}


def normalized_score_key(value: Any) -> str:
    return str(value).strip().strip("'\"")


def to_float_score(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        score = float(value)
    except (TypeError, ValueError):
        return None
    return score if math.isfinite(score) else None


def normalized_scores(row: dict[str, Any]) -> dict[str, float]:
    scores = row.get("scores") or {}
    normalized = {}
    for key, value in scores.items():
        score = to_float_score(value)
        if score is not None:
            normalized[normalized_score_key(key)] = score
    return normalized


def slugify(value: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9._-]+", "_", value.strip()).strip("_")
    return slug or "model"


def resolve_result_dir(results_dir: Path, result_id: str, *, label: str) -> Path:
    result_dir = results_dir / result_id
    if not result_dir.exists():
        raise FileNotFoundError(f"{label} result folder not found: {result_dir}")
    return result_dir


def resolve_specified_scores_path(result_dir: Path) -> Path:
    path = result_dir / "specified_metric_scores.jsonl"
    if not path.exists():
        raise FileNotFoundError(f"Specified metrics score file not found: {path}")
    return path


def build_rows(
    tasks: dict[str, dict[str, Any]],
    evaluation_records: list[dict[str, Any]],
    specified_records: list[dict[str, Any]],
) -> list[list[str | float | None]]:
    eval_by_task = {row.get("task_id"): row for row in evaluation_records}
    specified_by_task = {row.get("task_id"): row for row in specified_records}
    rows: list[list[str | float | None]] = []

    for task_id in sorted(tasks):
        task = tasks[task_id]
        task_type = str(task.get("task", {}).get("type") or "")
        eval_scores = normalized_scores(eval_by_task.get(task_id, {}))
        specified_scores = normalized_scores(specified_by_task.get(task_id, {}))

        row: list[str | float | None] = [
            task_id,
            TASK_TYPE_ZH.get(task_type, task_type),
            str((task.get("video") or {}).get("id") or ""),
            str((task.get("audio") or {}).get("id") or ""),
        ]
        row.extend(eval_scores.get(metric) for metric in GENERAL_METRICS)

        allowed_specified_metrics = set(SPECIFIED_METRICS_BY_TYPE.get(task_type, []))
        for metric in SPECIFIED_METRICS:
            row.append(specified_scores.get(metric) if metric in allowed_specified_metrics else None)
        rows.append(row)
    return rows


def style_workbook(ws, *, max_row: int, max_col: int) -> None:
    title_fill = PatternFill("solid", fgColor="E7EEF8")
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    base_fill = PatternFill("solid", fgColor="EEF5FF")
    general_fill = PatternFill("solid", fgColor="E8F4E8")
    specified_fill = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="B7C9D8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for cell in ws[2]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for col_idx in range(1, 5):
        ws.cell(row=2, column=col_idx).fill = base_fill
    for col_idx in range(5, 12):
        ws.cell(row=2, column=col_idx).fill = general_fill
    for col_idx in range(12, max_col + 1):
        ws.cell(row=2, column=col_idx).fill = specified_fill

    for row_idx in range(3, max_row + 1):
        ws.cell(row=row_idx, column=1).font = Font(bold=True)
        for col_idx in range(5, max_col + 1):
            ws.cell(row=row_idx, column=col_idx).number_format = "0.00"

    widths = [14, 10, 12, 12] + [12] * (max_col - 4)
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 24
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = ws.dimensions


def write_xlsx(path: Path, model_name: str, rows: list[list[str | float | None]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Details"

    headers = ["任务ID", "类型", "视频", "音频", *GENERAL_METRICS, *SPECIFIED_METRICS]
    ws["A1"] = model_name
    ws.append(headers)
    for row in rows:
        ws.append(row)

    style_workbook(ws, max_row=len(rows) + 2, max_col=len(headers))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Export per-task evaluation details to Excel.")
    parser.add_argument("--model-name", required=True, help="Model or method name shown in the first row.")
    parser.add_argument("--eval-id", required=True, help="General evaluation result folder name under eval_results/.")
    parser.add_argument("--specified-metrics-id", required=True, help="Specified metrics result folder name under eval_results/.")
    parser.add_argument("--results-dir", default=str(DEFAULT_RESULTS_DIR), help="Directory containing evaluation result folders.")
    parser.add_argument(
        "--output",
        default=None,
        help="Output .xlsx path. Defaults to eval_results/<eval_id>/<eval_id>_<model>_details.xlsx.",
    )
    args = parser.parse_args()

    results_dir = Path(args.results_dir)
    if not results_dir.is_absolute():
        results_dir = ROOT / results_dir

    eval_dir = resolve_result_dir(results_dir, args.eval_id, label="Evaluation")
    specified_dir = resolve_result_dir(results_dir, args.specified_metrics_id, label="Specified metrics")
    evaluation_scores_path = eval_dir / "evaluation_scores.jsonl"
    if not evaluation_scores_path.exists():
        raise FileNotFoundError(f"Evaluation scores file not found: {evaluation_scores_path}")
    specified_scores_path = resolve_specified_scores_path(specified_dir)

    rows = build_rows(
        load_tasks(),
        load_jsonl(evaluation_scores_path),
        load_jsonl(specified_scores_path),
    )

    output_path = (
        Path(args.output)
        if args.output
        else eval_dir / f"{args.eval_id}_{slugify(args.model_name)}_details.xlsx"
    )
    if not output_path.is_absolute():
        output_path = ROOT / output_path
    write_xlsx(output_path, args.model_name, rows)
    print(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
